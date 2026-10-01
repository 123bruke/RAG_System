"""Document loaders: PDF / TXT / DOCX  ->  text + metadata.

Design: one small class per file type, all sharing the :class:`BaseLoader`
interface, plus a :class:`DocumentLoader` that picks the right one from the
file extension. Adding a new format (e.g. Markdown) means adding one class and
one registry entry - nothing else changes.

Every loader returns *pages*. PDFs have real page numbers; TXT and DOCX come
back as a single page with ``page=None`` (they have no fixed pagination).
"""

import logging
from abc import ABC, abstractmethod
from dataclasses import dataclass
from pathlib import Path

from app.exceptions import (
    CorruptedDocumentError,
    DocumentNotFoundError,
    EmptyDocumentError,
    RAGError,
    UnsupportedFileError,
)

logger = logging.getLogger(__name__)


@dataclass
class DocumentPage:
    """Text of one page (``page`` is None when the format has no pages)."""

    text: str
    page: int | None = None


@dataclass
class LoadedDocument:
    """A parsed file: where it came from, its type and its pages of text."""

    source: str  # filename, e.g. "refund_policy.pdf"
    doc_type: str  # "pdf" | "txt" | "docx"
    pages: list[DocumentPage]


class BaseLoader(ABC):
    """Interface every file-type loader implements."""

    doc_type: str

    @abstractmethod
    def load(self, path: Path) -> list[DocumentPage]:
        """Read ``path`` and return its pages of text."""


class PDFLoader(BaseLoader):
    """Extract text page-by-page with pypdf (text PDFs only - no OCR)."""

    doc_type = "pdf"

    def load(self, path: Path) -> list[DocumentPage]:
        from pypdf import PdfReader

        try:
            reader = PdfReader(str(path))
            if reader.is_encrypted and not reader.decrypt(""):
                raise CorruptedDocumentError(
                    f"'{path.name}' is password-protected and cannot be read."
                )
            return [
                DocumentPage(text=page.extract_text() or "", page=number)
                for number, page in enumerate(reader.pages, start=1)
            ]
        except RAGError:
            raise
        except Exception as exc:  # pypdf raises many different error types
            raise CorruptedDocumentError(f"Could not read PDF '{path.name}': {exc}") from exc


class TxtLoader(BaseLoader):
    """Read plain text (UTF-8, falling back to Latin-1 for legacy files)."""

    doc_type = "txt"

    def load(self, path: Path) -> list[DocumentPage]:
        raw = path.read_bytes()
        try:
            text = raw.decode("utf-8-sig")
        except UnicodeDecodeError:
            logger.warning("%s is not valid UTF-8; decoding as Latin-1", path.name)
            text = raw.decode("latin-1")
        return [DocumentPage(text=text, page=None)]


class DocxLoader(BaseLoader):
    """Extract paragraphs and table cells from a Word document.

    Limitation: tables are appended after the paragraphs, so their position
    relative to the surrounding text is not preserved.
    """

    doc_type = "docx"

    def load(self, path: Path) -> list[DocumentPage]:
        from docx import Document

        try:
            document = Document(str(path))
            parts = [p.text for p in document.paragraphs if p.text.strip()]
            for table in document.tables:
                for row in table.rows:
                    cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
        except Exception as exc:  # BadZipFile, PackageNotFoundError, KeyError...
            raise CorruptedDocumentError(f"Could not read DOCX '{path.name}': {exc}") from exc
        return [DocumentPage(text="\n".join(parts), page=None)]


class DocumentLoader:
    """Facade: choose a loader by extension, validate, and return a LoadedDocument."""

    _LOADERS: dict[str, BaseLoader] = {
        ".pdf": PDFLoader(),
        ".txt": TxtLoader(),
        ".docx": DocxLoader(),
    }

    @classmethod
    def supported_extensions(cls) -> list[str]:
        return sorted(cls._LOADERS)

    @classmethod
    def is_supported(cls, path: str | Path) -> bool:
        return Path(path).suffix.lower() in cls._LOADERS

    def load(self, path: str | Path) -> LoadedDocument:
        path = Path(path)
        if not path.is_file():
            raise DocumentNotFoundError(f"File not found: {path}")

        extension = path.suffix.lower()
        loader = self._LOADERS.get(extension)
        if loader is None:
            supported = ", ".join(self.supported_extensions())
            raise UnsupportedFileError(
                f"Unsupported file type '{extension or '(none)'}' for '{path.name}'. "
                f"Supported types: {supported}."
            )

        logger.info("Loading document: %s", path.name)
        pages = [p for p in loader.load(path) if p.text.strip()]
        if not pages:
            raise EmptyDocumentError(
                f"'{path.name}' contains no extractable text "
                "(scanned PDFs need OCR, which this project does not include)."
            )
        return LoadedDocument(source=path.name, doc_type=loader.doc_type, pages=pages)
