from pathlib import Path

import pytest
from docx import Document

from app.config import Settings
from app.exceptions import LLMAuthError, NoDocumentsError
from app.generation.llm import create_llm
from app.generation.prompts import NO_ANSWER_MESSAGE
from app.pipeline import RAGPipeline, extract_citations
from app.schemas import HybridResult


def test_ingest_stores_chunks_and_builds_bm25(pipeline, sample_files):
    report = pipeline.ingest_documents(sample_files)
    assert report.ok_count == 3 and report.failed_count == 0
    assert report.total_chunks == pipeline.vector_store.count() > 3
    assert pipeline.keyword_index.count == report.total_chunks


def test_query_returns_answer_and_sources(pipeline, sample_files, llm):
    pipeline.ingest_documents(sample_files)
    response = pipeline.ask("What is the refund period?")

    assert response.answer == "The refund period is 30 days. [1]"
    assert response.answered is True
    assert response.cited == {1}
    assert 1 <= len(response.sources) <= 3  # TOP_K_RERANK
    assert response.sources == response.retrieval.reranked  # sources ARE the reranked context
    assert all(src.rerank_score is not None for src in response.sources)
    scores = [src.rerank_score for src in response.sources]
    assert scores == sorted(scores, reverse=True)
    assert "refund_policy.txt" in {src.source for src in response.sources}
    # the LLM received the question, numbered context and source metadata
    system_prompt, user_prompt = llm.calls[0]
    assert "I don't know based on the provided documents." in system_prompt
    assert "What is the refund period?" in user_prompt
    for number, src in enumerate(response.sources, start=1):  # numbered, in order, with metadata
        assert f"[{number}] Source: {src.source} | chunk {src.chunk_index}" in user_prompt
        assert src.text in user_prompt


def test_response_exposes_every_retrieval_stage(pipeline, sample_files):
    pipeline.ingest_documents(sample_files)
    r = pipeline.ask("HP:0001945").retrieval
    assert r.vector_hits and r.bm25_hits and r.fused and r.reranked
    assert r.reranked[0].source == "hpo_phenotype_terms.txt"
    assert "HP:0001945 Fever" in r.reranked[0].text


def test_top_k_overrides(pipeline, sample_files):
    pipeline.ingest_documents(sample_files)
    r = pipeline.ask("refund", top_k_vector=2, top_k_bm25=2, top_k_hybrid=3, top_k_rerank=1).retrieval
    assert len(r.vector_hits) == 2 and len(r.bm25_hits) <= 2
    assert len(r.fused) <= 3 and len(r.reranked) == 1


def test_query_without_documents_raises(pipeline):
    with pytest.raises(NoDocumentsError):
        pipeline.ask("Anything?")


def test_no_result_behavior_skips_the_llm(pipeline, sample_files, llm, monkeypatch):
    pipeline.ingest_documents(sample_files)
    monkeypatch.setattr(pipeline, "retrieve", lambda *a, **k: HybridResult())
    response = pipeline.ask("Something unrelated")
    assert response.answer == NO_ANSWER_MESSAGE
    assert response.answered is False and response.sources == []
    assert llm.calls == []  # never called the LLM


def test_llm_saying_it_does_not_know_is_flagged(pipeline, sample_files, llm):
    pipeline.ingest_documents(sample_files)
    llm.answer = NO_ANSWER_MESSAGE
    assert pipeline.ask("What colour is the moon?").answered is False


def test_empty_question_is_rejected(pipeline, sample_files):
    pipeline.ingest_documents(sample_files)
    with pytest.raises(ValueError):
        pipeline.ask("   ")


def test_reingesting_replaces_instead_of_duplicating(pipeline, sample_files):
    pipeline.ingest_documents(sample_files)
    first = pipeline.vector_store.count()
    pipeline.ingest_documents(sample_files)
    assert pipeline.vector_store.count() == first
    assert pipeline.keyword_index.count == first


def test_bad_files_are_reported_and_do_not_block_good_ones(pipeline, sample_files, tmp_path):
    unsupported = tmp_path / "notes.xyz"
    unsupported.write_text("data")
    empty = tmp_path / "empty.txt"
    empty.write_text("   \n")
    corrupt = tmp_path / "broken.pdf"
    corrupt.write_bytes(b"this is definitely not a pdf")
    missing = tmp_path / "ghost.txt"

    report = pipeline.ingest_documents([unsupported, empty, corrupt, missing, sample_files[0]])
    by_name = {f.filename: f for f in report.files}
    assert by_name["notes.xyz"].status == "error" and "Unsupported" in by_name["notes.xyz"].error
    assert by_name["empty.txt"].status == "error" and "no extractable text" in by_name["empty.txt"].error
    assert by_name["broken.pdf"].status == "error" and "Could not read PDF" in by_name["broken.pdf"].error
    assert by_name["ghost.txt"].status == "error" and "not found" in by_name["ghost.txt"].error
    assert by_name[sample_files[0].name].status == "ok"
    assert pipeline.vector_store.count() > 0


def test_docx_ingestion(pipeline, tmp_path):
    path = tmp_path / "handbook.docx"
    doc = Document()
    doc.add_paragraph("Employees receive 25 days of paid vacation per year.")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Notice period"
    table.rows[0].cells[1].text = "2 months"
    doc.save(path)

    report = pipeline.ingest_documents([path])
    assert report.ok_count == 1
    text = " ".join(c.text for c in pipeline.vector_store.get_all())
    assert "25 days of paid vacation" in text and "Notice period | 2 months" in text
    assert pipeline.vector_store.get_all()[0].metadata["doc_type"] == "docx"


def test_clear_empties_everything(pipeline, sample_files):
    pipeline.ingest_documents(sample_files)
    pipeline.clear()
    assert pipeline.vector_store.count() == 0 and pipeline.keyword_index.count == 0
    with pytest.raises(NoDocumentsError):
        pipeline.ask("refund?")


def test_index_survives_restart_and_bm25_is_rebuilt(pipeline, settings, sample_files):
    pipeline.ingest_documents(sample_files)
    restarted = RAGPipeline(
        settings, embedder=pipeline.embedder, reranker=pipeline.reranker, llm=pipeline.llm
    )
    response = restarted.ask("HP:0001945")  # BM25 must be rebuilt from ChromaDB
    assert restarted.keyword_index.count > 0
    assert response.retrieval.bm25_hits


def test_stats(pipeline, sample_files):
    pipeline.ingest_documents(sample_files)
    stats = pipeline.stats()
    assert stats["chunks"] == pipeline.vector_store.count()
    assert set(stats["documents"]) == {"refund_policy.txt", "customer_support_policy.txt", "hpo_phenotype_terms.txt"}
    assert stats["top_k"]["rerank"] == 3


def test_ingest_directory(pipeline, settings, sample_files):
    docs_dir = Path(settings.documents_dir)
    docs_dir.mkdir(parents=True)
    for f in sample_files:
        (docs_dir / f.name).write_text(f.read_text(encoding="utf-8"), encoding="utf-8")
    (docs_dir / "ignore.me").write_text("x")
    assert pipeline.ingest_directory().ok_count == 3


# ---------------------------------------------------------- small helpers
@pytest.mark.parametrize(
    "answer,n,expected",
    [
        ("Yes [1].", 3, {1}),
        ("See [1][3] and [2, 3].", 3, {1, 2, 3}),
        ("Made up [9].", 3, set()),
        ("No citations.", 3, set()),
    ],
)
def test_extract_citations(answer, n, expected):
    assert extract_citations(answer, n) == expected


def test_create_llm_without_key_raises_clear_error():
    with pytest.raises(LLMAuthError, match="GOOGLE_API_KEY"):
        create_llm(Settings(_env_file=None, llm_provider="gemini", google_api_key=""))
    with pytest.raises(LLMAuthError, match="OPENAI_API_KEY"):
        create_llm(Settings(_env_file=None, llm_provider="openai", openai_api_key=""))


def test_settings_validation_and_defaults():
    with pytest.raises(ValueError, match="CHUNK_OVERLAP"):
        Settings(_env_file=None, chunk_size=100, chunk_overlap=100)
    s = Settings(_env_file=None, llm_provider=" OpenAI ")
    assert s.llm_provider == "openai" and s.resolved_llm_model == "gpt-4o-mini"
    assert Settings(_env_file=None).resolved_llm_model == "gemini-2.5-flash"
